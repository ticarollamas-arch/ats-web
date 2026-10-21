"""Auditoria de arquivo (.pdf / .docx) para problemas de parsing ATS."""
from __future__ import annotations

import io
import re

import resume_tools as rt

HEADINGS = {
    "Resumo / Summary": r"\b(resumo( profissional)?|sobre mim|perfil( profissional)?|summary|profile|objetivo|objective)\b",
    "Experiência / Experience": r"\b(experi[eê]ncias?( profissionais?)?|hist[oó]rico profissional|work experience|experience|employment)\b",
    "Formação / Education": r"\b(forma[cç][aã]o( acad[eê]mica)?|educa[cç][aã]o|escolaridade|education)\b",
    "Competências / Skills": r"\b(compet[eê]ncias|habilidades|skills|conhecimentos|tecnologias|technical skills)\b",
}
MONTHS = r"(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez|feb|apr|may|aug|sep|oct|dec|janeiro|fevereiro|mar[cç]o|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro|january|february|march|april|june|july|august|september|october|november|december)"
DATE_RE = re.compile(rf"\b({MONTHS}\.?[\s/\-]*(de\s+)?\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})\s*[–\-—a]+\s*(atual|presente|present|{MONTHS}\.?[\s/\-]*(de\s+)?\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})", re.I)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d{1,3}[\s.-]?)?\(?\d{2,3}\)?[\s.-]?9?\d{3,4}[\s.-]?\d{4}")


def _item(level: str, code: str, msg: str, fix: str = "") -> dict:
    return {"level": level, "code": code, "message": msg, "fix": fix}


def _docx_structure(data: bytes) -> tuple[str, list[dict]]:
    from docx import Document
    from docx.oxml.ns import qn

    items: list[dict] = []
    doc = Document(io.BytesIO(data))
    body = doc.element.body
    text_parts = [p.text for p in doc.paragraphs]

    if doc.tables:
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    text_parts.append(cell.text)
        items.append(_item("error", "tables", f"{len(doc.tables)} tabela(s) encontrada(s). Muitos ATS leem tabelas fora de ordem ou ignoram o conteúdo.",
                           "Troque tabelas por parágrafos e marcadores simples, em coluna única."))
    txbx = body.findall(".//" + qn("w:txbxContent"))
    if txbx:
        items.append(_item("error", "textbox", f"{len(txbx)} caixa(s) de texto. Texto dentro de caixas costuma ficar invisível ao parser.",
                           "Mova o conteúdo para o corpo normal do documento."))
    for sec in doc.sections:
        cols = sec._sectPr.find(qn("w:cols"))
        if cols is not None and int(cols.get(qn("w:num"), "1")) > 1:
            items.append(_item("error", "columns", "Layout de múltiplas colunas detectado.", "Use uma única coluna."))
            break
    hf = []
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            t = " ".join(p.text for p in part.paragraphs).strip()
            if t:
                hf.append(t)
    if hf:
        items.append(_item("error", "header_footer", "Há texto em cabeçalho/rodapé (" + hf[0][:60] + "…). Vários ATS ignoram essas áreas.",
                           "Coloque nome e contato no corpo do documento."))
    imgs = len(doc.inline_shapes) + len(body.findall(".//" + qn("w:drawing"))) // 2
    if imgs:
        items.append(_item("warn", "images", "Imagens/gráficos/ícones detectados. Não são lidos pelo ATS e podem quebrar o parsing.",
                           "Remova fotos, ícones, barras de nível de habilidade e logos."))
    return "\n".join(text_parts), items


def _pdf_structure(data: bytes) -> tuple[str, int, list[dict]]:
    from pypdf import PdfReader

    items: list[dict] = []
    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
    if len(text.strip()) < 100:
        items.append(_item("error", "no_text", "Quase nenhum texto selecionável. O PDF parece ser imagem/escaneado: o ATS verá um documento em branco.",
                           "Exporte o PDF a partir do Word/Docs/Google Docs (texto real), não de scanner ou captura de tela."))
    imgs = sum(len(pg.images) for pg in reader.pages) if text.strip() else 0
    if imgs:
        items.append(_item("warn", "images", f"{imgs} imagem(ns) no PDF (foto, ícones ou logos).", "Remova imagens decorativas."))
    # heurística de colunas: linhas muito curtas e muitas
    lines = [l for l in text.splitlines() if l.strip()]
    if lines:
        short = sum(1 for l in lines if len(l.strip()) < 25)
        if len(lines) > 25 and short / len(lines) > 0.55:
            items.append(_item("warn", "columns_hint", "Muitas linhas curtas: possível layout em colunas ou texto quebrado na extração.",
                               "Confira se o currículo é de coluna única; abra o texto extraído abaixo e veja se a ordem faz sentido."))
    return text, pages, items


def audit_file(filename: str, data: bytes, jd: str = "") -> dict:
    name = (filename or "").lower()
    pages = None
    if name.endswith(".docx"):
        kind = "docx"
        text, items = _docx_structure(data)
    elif name.endswith(".pdf"):
        kind = "pdf"
        text, pages, items = _pdf_structure(data)
    else:
        raise ValueError("Formato não suportado. Envie .pdf ou .docx.")

    low = text.lower()
    nonempty = bool(text.strip())

    if nonempty:
        if not EMAIL_RE.search(text):
            items.append(_item("error", "email", "Nenhum e-mail encontrado no texto.", "Inclua e-mail no corpo do currículo."))
        if not PHONE_RE.search(text):
            items.append(_item("warn", "phone", "Nenhum telefone reconhecível.", "Inclua telefone com DDD, ex.: (11) 91234-5678."))
        found = [k for k, rx in HEADINGS.items() if re.search(rx, low)]
        missing = [k for k in HEADINGS if k not in found]
        if missing:
            items.append(_item("warn", "headings", "Títulos de seção padrão não encontrados: " + ", ".join(missing) + ".",
                               "Use títulos convencionais (Experiência Profissional, Formação Acadêmica, Competências). Títulos criativos confundem o parser."))
        dates = DATE_RE.findall(text)
        if not dates:
            items.append(_item("warn", "dates", "Nenhum intervalo de datas reconhecido.", "Use o formato 'Mar 2021 – Atual' em cada experiência."))
        if re.search(r"[■●◆►★☆✔✓➤❖]|[]", text):
            items.append(_item("warn", "bullets", "Símbolos de marcador não padrão detectados.", "Use marcadores simples (•) do próprio editor."))
        words = len(re.findall(r"\w+", text))
        if pages and pages > 2:
            items.append(_item("warn", "length", f"{pages} páginas. Para a maioria dos perfis, 1 página (até ~10 anos) ou 2 páginas é o ideal.", "Corte o que não apoia a vaga-alvo."))
        if words < 150:
            items.append(_item("warn", "short", f"Apenas {words} palavras: pouco conteúdo para o parser e o recrutador.", "Detalhe resultados mensuráveis por cargo."))
        nums = len(re.findall(r"\d+[%.,]?\d*", text))
        if nums < 6:
            items.append(_item("warn", "metrics", "Poucos números no texto. Currículos fortes quantificam resultados (%, R$, volume, tempo).",
                               "Veja a aba Guia > Metric proxies, sem inventar números."))

    errors = sum(1 for i in items if i["level"] == "error")
    warns = sum(1 for i in items if i["level"] == "warn")
    if not items:
        items.append(_item("ok", "clean", "Nenhum problema estrutural detectado.", ""))

    out = {
        "kind": kind,
        "pages": pages,
        "words": len(re.findall(r"\w+", text)),
        "errors": errors,
        "warnings": warns,
        "verdict": "problemas críticos de parsing" if errors else ("ajustes recomendados" if warns else "estrutura segura para ATS"),
        "items": items,
        "extracted_text": text[:20000],
        "note": "Auditoria estrutural heurística. Não é a nota interna de nenhum ATS real.",
    }

    if jd.strip() and nonempty:
        req, pref = rt.extract_keywords(jd)
        rp, rm = rt.coverage_of(req, low)
        pp, pm = rt.coverage_of(pref, low)
        out["coverage"] = {
            "required_pct": round(100 * len(rp) / len(req), 1) if req else 100.0,
            "preferred_pct": round(100 * len(pp) / len(pref), 1) if pref else 100.0,
            "missing_required": sorted(rm),
            "missing_preferred": sorted(pm),
        }
    return out
